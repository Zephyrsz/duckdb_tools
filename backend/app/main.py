from __future__ import annotations

import csv
import re
import time
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

import duckdb
import openpyxl
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from .duckdb_manager import DuckDBBusyError, DuckDBConnectionManager, DuckDBNotConnectedError
from .semantic import SemanticRepository, build_drafts, profile_database


PROJECT_ROOT = Path(__file__).resolve().parents[2]
RDS_AGENT_ROOT = Path(__import__("os").environ.get("RDS_AGENT_ROOT", PROJECT_ROOT.parent / "rds_agent"))
DATA_DIR = Path(__import__("os").environ.get("DUCKDB_TOOLS_DATA", PROJECT_ROOT / "data"))
UPLOAD_DIR = DATA_DIR / "uploads"
DB_PATH = Path(__import__("os").environ.get("DUCKDB_TOOLS_DATABASE", DATA_DIR / "workspace.duckdb"))
METADATA_DB_PATH = Path(__import__("os").environ.get("RDS_AGENT_METADATA_DB", RDS_AGENT_ROOT / "var" / "metadata.db"))
MAX_UPLOAD_BYTES = 50 * 1024 * 1024
PREVIEW_ROWS = 5
TABLE_ROWS = 100
ALLOWED_EXTENSIONS = {".csv": "csv", ".xlsx": "excel"}
TABLE_NAME_PATTERN = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
READ_ONLY_KEYWORDS = {"SELECT", "WITH", "DESCRIBE", "SHOW", "EXPLAIN"}
FORBIDDEN_SQL = re.compile(
    r"\b(INSERT|UPDATE|DELETE|CREATE|DROP|ALTER|COPY|ATTACH|DETACH|INSTALL|LOAD|EXPORT|PRAGMA|VACUUM|CALL)\b",
    re.IGNORECASE,
)
FORBIDDEN_SCANS = re.compile(r"\b(read_csv|read_json|read_parquet|parquet_scan|glob|httpfs)\b", re.IGNORECASE)


app = FastAPI(title="DuckDB Tools API", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["GET", "POST", "PATCH"],
    allow_headers=["*"],
)


class ImportRequest(BaseModel):
    stored_path: str
    table_name: str = Field(min_length=1, max_length=96)
    has_header: bool = True


class QueryRequest(BaseModel):
    sql: str = Field(min_length=1, max_length=100_000)


class SemanticDraftPatch(BaseModel):
    canonical_name: str | None = Field(default=None, min_length=1, max_length=96)
    display_name: str | None = Field(default=None, min_length=1, max_length=160)
    description: str | None = Field(default=None, max_length=500)
    expression: str | None = Field(default=None, max_length=2_000)
    status: str | None = None
    confidence: float | None = Field(default=None, ge=0, le=1)


def ensure_storage() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    resolve_database_path().parent.mkdir(parents=True, exist_ok=True)
    SemanticRepository(METADATA_DB_PATH)
    duckdb_manager.initialize()


def resolve_database_path() -> Path:
    """Return the configured DuckDB file, preserving the test/runtime override."""
    configured = __import__("os").environ.get("DUCKDB_TOOLS_DATABASE")
    return Path(configured) if configured else DB_PATH


duckdb_manager = DuckDBConnectionManager(resolve_database_path)


def connect_db() -> duckdb.DuckDBPyConnection:
    ensure_storage()
    return duckdb_manager.lease()  # type: ignore[return-value]


def semantic_repository() -> SemanticRepository:
    ensure_storage()
    return SemanticRepository(METADATA_DB_PATH)


def semantic_profiles() -> dict[str, dict[str, Any]]:
    with duckdb_manager.operation() as connection:
        return {profile["name"]: profile for profile in profile_database(connection)}


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if hasattr(value, "as_tuple"):
        return float(value)
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return value


def rows_as_dicts(description: list[tuple[Any, ...]], rows: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    names = [str(column[0]) for column in description]
    return [{name: json_value(value) for name, value in zip(names, row)} for row in rows]


def columns_from_description(description: list[tuple[Any, ...]]) -> list[dict[str, str]]:
    return [{"name": str(column[0]), "type": str(column[1])} for column in description]


def normalize_headers(headers: list[Any]) -> list[str]:
    used: set[str] = set()
    normalized: list[str] = []
    for index, value in enumerate(headers):
        base = str(value).strip() if value not in (None, "") else f"column_{index + 1}"
        base = re.sub(r"\s+", "_", base)
        candidate = base
        suffix = 2
        while candidate in used:
            candidate = f"{base}_{suffix}"
            suffix += 1
        used.add(candidate)
        normalized.append(candidate)
    return normalized


def infer_excel_type(values: list[Any]) -> str:
    present = [value for value in values if value is not None and value != ""]
    if not present:
        return "VARCHAR"
    if all(isinstance(value, bool) for value in present):
        return "BOOLEAN"
    if all(isinstance(value, int) and not isinstance(value, bool) for value in present):
        return "BIGINT"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in present):
        return "DOUBLE"
    if all(isinstance(value, datetime) for value in present):
        return "TIMESTAMP"
    if all(isinstance(value, date) for value in present):
        return "DATE"
    return "VARCHAR"


def parse_csv_value(value: str) -> Any:
    value = value.strip()
    if not value:
        return None
    lowered = value.lower()
    if lowered in {"true", "false"}:
        return lowered == "true"
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


def read_csv_preview(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        raw_rows = list(csv.reader(stream))
    if not raw_rows:
        raise ValueError("CSV 文件没有可读取的表头")
    headers = normalize_headers(raw_rows[0])
    body = [
        [parse_csv_value(value) for value in row[: len(headers)]]
        + [None] * max(0, len(headers) - len(row))
        for row in raw_rows[1:]
    ]
    columns = [{"name": name, "type": infer_excel_type([row[index] for row in body])} for index, name in enumerate(headers)]
    rows = [dict(zip(headers, row)) for row in body[:PREVIEW_ROWS]]
    return rows, columns


def read_excel(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, str]], list[list[Any]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not raw_rows:
        raise ValueError("Excel 文件没有可读取的表头")
    headers = normalize_headers(raw_rows[0])
    body = [row[: len(headers)] + [None] * max(0, len(headers) - len(row)) for row in raw_rows[1:]]
    columns = [{"name": name, "type": infer_excel_type([row[index] for row in body])} for index, name in enumerate(headers)]
    rows = [{name: json_value(value) for name, value in zip(headers, row)} for row in body[:PREVIEW_ROWS]]
    return rows, columns, body


def preview_upload(path: Path, kind: str) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    if kind == "excel":
        rows, columns, _ = read_excel(path)
        return rows, columns
    return read_csv_preview(path)


def kind_for_filename(filename: str) -> str:
    extension = Path(filename).suffix.lower()
    kind = ALLOWED_EXTENSIONS.get(extension)
    if not kind:
        raise HTTPException(status_code=400, detail="仅支持 CSV 或 XLSX 文件")
    return kind


def safe_upload_path(stored_path: str) -> Path:
    path = Path(stored_path).resolve()
    upload_root = UPLOAD_DIR.resolve()
    if upload_root not in path.parents:
        raise HTTPException(status_code=400, detail="上传文件路径无效")
    if not path.exists() or not path.is_file():
        raise HTTPException(status_code=404, detail="上传文件不存在，请重新上传")
    return path


def quote_identifier(identifier: str) -> str:
    if not TABLE_NAME_PATTERN.fullmatch(identifier):
        raise HTTPException(status_code=400, detail="表名只能包含字母、数字和下划线，且不能以数字开头")
    return '"' + identifier.replace('"', '""') + '"'


def table_columns(connection: duckdb.DuckDBPyConnection, table_name: str) -> list[dict[str, str]]:
    rows = connection.execute(f"DESCRIBE {quote_identifier(table_name)}").fetchall()
    return [{"name": str(row[0]), "type": str(row[1])} for row in rows]


def database_not_connected(error: DuckDBNotConnectedError) -> HTTPException:
    return HTTPException(status_code=503, detail=str(error))


def database_busy(error: DuckDBBusyError) -> HTTPException:
    return HTTPException(status_code=409, detail=str(error))


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/api/duckdb/status")
def duckdb_status() -> dict[str, object]:
    return duckdb_manager.status()


@app.post("/api/duckdb/connect")
def duckdb_connect() -> dict[str, object]:
    try:
        return duckdb_manager.connect()
    except DuckDBBusyError as error:
        raise database_busy(error) from error


@app.post("/api/duckdb/disconnect")
def duckdb_disconnect() -> dict[str, object]:
    try:
        return duckdb_manager.disconnect()
    except DuckDBBusyError as error:
        raise database_busy(error) from error


@app.post("/api/upload/preview")
async def upload_preview(file: UploadFile = File(...)) -> dict[str, Any]:
    filename = Path(file.filename or "upload").name
    kind = kind_for_filename(filename)
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="文件不能超过 50 MB")
    ensure_storage()
    stored_path = UPLOAD_DIR / f"{uuid4().hex}{Path(filename).suffix.lower()}"
    stored_path.write_bytes(content)
    try:
        rows, columns = preview_upload(stored_path, kind)
    except Exception as error:
        stored_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=f"文件解析失败：{error}") from error
    return {"filename": filename, "stored_path": str(stored_path), "kind": kind, "columns": columns, "preview": rows}


@app.post("/api/import")
def import_file(request: ImportRequest) -> dict[str, Any]:
    path = safe_upload_path(request.stored_path)
    table_name = request.table_name.strip()
    quoted_table = quote_identifier(table_name)
    kind = ALLOWED_EXTENSIONS.get(path.suffix.lower())
    if not kind:
        raise HTTPException(status_code=400, detail="上传文件格式无效")
    try:
        with duckdb_manager.operation() as connection:
            connection.execute("BEGIN")
            try:
                if kind == "csv":
                    connection.execute(
                        f"CREATE OR REPLACE TABLE {quoted_table} AS SELECT * FROM read_csv_auto(?, header = ?)",
                        [str(path), request.has_header],
                    )
                else:
                    _, columns, body = read_excel(path)
                    if not columns:
                        raise HTTPException(status_code=400, detail="Excel 文件没有可导入的列")
                    column_sql = ", ".join(f"{quote_identifier(column['name'])} {column['type']}" for column in columns)
                    connection.execute(f"CREATE OR REPLACE TABLE {quoted_table} ({column_sql})")
                    placeholders = ", ".join("?" for _ in columns)
                    connection.executemany(f"INSERT INTO {quoted_table} VALUES ({placeholders})", body)
                connection.execute("COMMIT")
            except Exception:
                connection.execute("ROLLBACK")
                raise
            row_count = connection.execute(f"SELECT COUNT(*) FROM {quoted_table}").fetchone()[0]
            return {"table_name": table_name, "row_count": row_count, "columns": table_columns(connection, table_name)}
    except DuckDBNotConnectedError as error:
        raise database_not_connected(error) from error
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"导入失败：{error}") from error


@app.get("/api/database")
def database_info() -> dict[str, Any]:
    try:
        with duckdb_manager.operation() as connection:
            table_names = [str(row[0]) for row in connection.execute("SHOW TABLES").fetchall()]
            tables = []
            for name in table_names:
                count = connection.execute(f"SELECT COUNT(*) FROM {quote_identifier(name)}").fetchone()[0]
                tables.append({"name": name, "row_count": count})
            return {"database": resolve_database_path().name, "tables": tables}
    except DuckDBNotConnectedError as error:
        raise database_not_connected(error) from error


@app.get("/api/tables/{table_name}")
def table_info(table_name: str) -> dict[str, Any]:
    try:
        with duckdb_manager.operation() as connection:
            columns = table_columns(connection, table_name)
            result = connection.execute(f"SELECT * FROM {quote_identifier(table_name)} LIMIT {TABLE_ROWS}")
            rows = rows_as_dicts(result.description, result.fetchall())
            row_count = connection.execute(f"SELECT COUNT(*) FROM {quote_identifier(table_name)}").fetchone()[0]
            return {"table_name": table_name, "columns": columns, "rows": rows, "row_count": row_count}
    except duckdb.CatalogException as error:
        raise HTTPException(status_code=404, detail="数据表不存在") from error
    except DuckDBNotConnectedError as error:
        raise database_not_connected(error) from error


def validate_read_only_sql(sql: str) -> None:
    stripped = sql.strip()
    statement = stripped[:-1].rstrip() if stripped.endswith(";") else stripped
    if ";" in statement:
        raise HTTPException(status_code=400, detail="查询接口只允许执行一条只读 SQL")
    keyword_match = re.match(r"(?:^\s*(?:--[^\n]*\n|/\*.*?\*/\s*)*)([A-Za-z]+)", statement, re.DOTALL)
    keyword = keyword_match.group(1).upper() if keyword_match else ""
    if keyword not in READ_ONLY_KEYWORDS or FORBIDDEN_SQL.search(statement) or FORBIDDEN_SCANS.search(statement):
        raise HTTPException(status_code=400, detail="查询接口只允许执行 SELECT、WITH、DESCRIBE、SHOW 或 EXPLAIN 等只读 SQL")


@app.post("/api/query")
def run_query(request: QueryRequest) -> dict[str, Any]:
    validate_read_only_sql(request.sql)
    started = time.perf_counter()
    try:
        with duckdb_manager.operation() as connection:
            result = connection.execute(request.sql)
            rows = rows_as_dicts(result.description, result.fetchall()) if result.description else []
            return {
                "columns": columns_from_description(result.description or []),
                "rows": rows,
                "row_count": len(rows),
                "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
            }
    except duckdb.Error as error:
        raise HTTPException(status_code=400, detail=f"查询失败：{error}") from error
    except DuckDBNotConnectedError as error:
        raise database_not_connected(error) from error


@app.get("/api/semantic/summary")
def semantic_summary() -> dict[str, Any]:
    return semantic_repository().summary()


@app.get("/api/semantic/tables")
def semantic_tables() -> list[dict[str, Any]]:
    try:
        profiles = list(semantic_profiles().values())
        return semantic_repository().tables([{"name": profile["name"], "row_count": profile["row_count"]} for profile in profiles])
    except DuckDBNotConnectedError as error:
        raise database_not_connected(error) from error


@app.get("/api/semantic/tables/{table_name}")
def semantic_table(table_name: str) -> dict[str, Any]:
    if not TABLE_NAME_PATTERN.fullmatch(table_name):
        raise HTTPException(status_code=400, detail="表名无效")
    try:
        profile = semantic_profiles().get(table_name)
    except DuckDBNotConnectedError as error:
        raise database_not_connected(error) from error
    if profile is None:
        raise HTTPException(status_code=404, detail="数据表不存在")
    return semantic_repository().table(table_name, profile)


@app.post("/api/semantic/scan")
def scan_semantic() -> dict[str, Any]:
    try:
        with duckdb_manager.operation() as connection:
            profiles = profile_database(connection)
        drafts = semantic_repository().replace_drafts(build_drafts(profiles))
        return {"table_count": len(profiles), "draft_count": len(drafts), "drafts": drafts}
    except DuckDBNotConnectedError as error:
        raise database_not_connected(error) from error


@app.patch("/api/semantic/drafts/{draft_id}")
def update_semantic_draft(draft_id: int, request: SemanticDraftPatch) -> dict[str, Any]:
    try:
        result = semantic_repository().update_draft(draft_id, request.model_dump(exclude_none=True))
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    if result is None:
        raise HTTPException(status_code=404, detail="语义草稿不存在")
    return result


@app.post("/api/semantic/validate")
def validate_semantic() -> dict[str, Any]:
    try:
        return semantic_repository().validate(semantic_profiles())
    except DuckDBNotConnectedError as error:
        raise database_not_connected(error) from error


@app.post("/api/semantic/publish")
def publish_semantic() -> dict[str, Any]:
    try:
        return semantic_repository().publish(semantic_profiles())
    except DuckDBNotConnectedError as error:
        raise database_not_connected(error) from error
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error


ensure_storage()

FRONTEND_DIST = PROJECT_ROOT / "frontend" / "dist"
if FRONTEND_DIST.exists():
    app.mount("/", StaticFiles(directory=FRONTEND_DIST, html=True), name="frontend")
